import openpyxl as px


def Record(x, y, depthimg):
    wb = px.load_workbook("record.xlsx")
    ws = wb["Sheet1"]

    column_num = ws.max_column
    if column_num == 1 and ws.cell(row=1, column=column_num).value == None:
        column_num = 0
    print(column_num)
    for i in range(40):
        for j in range(40):
            cell = ws.cell(row=(i*40+(j+1)),column=column_num+1)
            cell.value = depthimg[y+j, x+i]
    
    wb.save("record.xlsx")